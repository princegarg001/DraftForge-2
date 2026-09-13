import io
import docx
from app.core.exceptions import ValidationFailedError
from app.core.logging import get_logger
from app.services.parsing.base_parser import BaseParser, ParsedDocument, ParsedPage
from app.utils.text_utils import clean_extracted_text

logger = get_logger("docx_parser")


class DocxParser(BaseParser):
    def parse(self, content: bytes, original_filename: str) -> ParsedDocument:
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [clean_extracted_text(p.text) for p in doc.paragraphs if p.text.strip()]
            
            # Extract content from tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([clean_extracted_text(cell.text) for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)

            combined_text = "\n\n".join(paragraphs)
            # DOCX does not have fixed layout pages, represented as a single continuous unit (Page 1)
            pages = [ParsedPage(page_number=1, text=combined_text)]

            return ParsedDocument(
                raw_text=combined_text,
                pages=pages,
                total_pages=1,
                metadata={"paragraphs_count": len(doc.paragraphs), "source_format": "docx"}
            )
        except Exception as exc:
            logger.error(f"DOCX extraction error on file '{original_filename}': {exc.__class__.__name__}: {exc}")
            raise ValidationFailedError(
                "This DOCX file could not be read. It may be corrupted or use an unsupported format."
            ) from exc